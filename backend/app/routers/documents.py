import os
import uuid
import logging

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, field_validator

from app.core.config import settings
from app.core.deps import require_role
from app.ai.llm_providers import call_llm
from app.services.java_bridge import bridge

logger = logging.getLogger("stratroom.documents")

router = APIRouter(prefix="/documents", tags=["documents"])

STORAGE_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "storage", "documents")
ALLOWED_EXTENSIONS = {"pdf", "docx", "xlsx"}
MAX_DOC_SIZE_MB = 20
os.makedirs(STORAGE_DIR, exist_ok=True)


def _safe_filename(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    safe = uuid.uuid4().hex[:16]
    return f"{safe}.{ext}" if ext else safe


def _validate_storage_path(path: str) -> str:
    real_storage = os.path.realpath(STORAGE_DIR)
    real_path = os.path.realpath(path)
    if not real_path.startswith(real_storage + os.sep) and real_path != real_storage:
        raise ValueError("Path traversal detected")
    return real_path


async def _extract_text(content: bytes, ext: str) -> str:
    if ext == "pdf":
        try:
            import fitz
            doc = fitz.open(stream=content, filetype="pdf")
            text_parts = [page.get_text() for page in doc]
            doc.close()
            return "\n".join(text_parts)[:50000]
        except Exception:
            return ""
    elif ext == "docx":
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)[:50000]
        except Exception:
            return ""
    elif ext == "xlsx":
        try:
            from openpyxl import load_workbook
            import io
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    parts.append(" | ".join(str(c) if c is not None else "" for c in row))
            wb.close()
            return "\n".join(parts)[:50000]
        except Exception:
            return ""
    return ""


@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    ctx: dict = Depends(require_role("manager")),
):
    user_id = ctx["user_id"]
    org_id = ctx["org_id"]
    uploaded_by = ctx.get("full_name", "")

    filename = file.filename or "document"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=422, detail=f"File type '.{ext}' not supported. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    content = await file.read()
    max_bytes = MAX_DOC_SIZE_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"File exceeds {MAX_DOC_SIZE_MB}MB limit")

    if not content:
        raise HTTPException(status_code=422, detail="Empty file")

    safe_name = _safe_filename(filename)
    storage_path = os.path.join(STORAGE_DIR, safe_name)
    _validate_storage_path(storage_path)

    with open(storage_path, "wb") as f:
        f.write(content)

    extracted = await _extract_text(content, ext)

    resp = await bridge.post(bridge.db_service, "/documents/upload", json={
        "org_id": org_id,
        "user_id": user_id,
        "filename": safe_name,
        "original_filename": filename,
        "content_type": file.content_type or f"application/{ext}",
        "file_size": len(content),
        "extracted_text": extracted,
        "storage_path": storage_path,
        "uploaded_by": uploaded_by,
    })
    doc_id = resp.get("id") if isinstance(resp, dict) else None

    logger.info("Document uploaded: id=%s user=%s file=%s size=%d", doc_id, user_id, filename, len(content))

    return {
        "id": doc_id,
        "filename": filename,
        "size": len(content),
        "content_type": file.content_type,
        "has_extracted_text": bool(extracted),
    }


@router.get("")
async def list_documents(
    ctx: dict = Depends(require_role("member")),
):
    result = await bridge.get(
        bridge.db_service, "/documents",
        params={
            "org_id": ctx["org_id"],
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    return {"documents": result if isinstance(result, list) else []}


@router.get("/search")
async def search_documents(
    q: str = Query(..., min_length=1, max_length=200),
    ctx: dict = Depends(require_role("member")),
):
    result = await bridge.get(
        bridge.db_service, "/documents/search",
        params={
            "q": q,
            "org_id": ctx["org_id"],
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    return {"documents": result if isinstance(result, list) else [], "query": q}


@router.get("/{document_id}")
async def get_document(
    document_id: int,
    ctx: dict = Depends(require_role("member")),
):
    result = await bridge.get(
        bridge.db_service, f"/documents/{document_id}",
        params={
            "org_id": ctx["org_id"],
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    if not result:
        raise HTTPException(status_code=404, detail="Document not found")

    summaries = await bridge.get(
        bridge.db_service, f"/documents/{document_id}/summaries",
    )
    doc = dict(result[0])
    doc["summaries"] = summaries if isinstance(summaries, list) else []
    return doc


@router.get("/{document_id}/download")
async def download_document(
    document_id: int,
    ctx: dict = Depends(require_role("member")),
):
    from fastapi.responses import FileResponse

    result = await bridge.get(
        bridge.db_service, f"/documents/{document_id}/download",
        params={
            "org_id": ctx["org_id"],
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    if not result:
        raise HTTPException(status_code=404, detail="Document not found")

    row = result[0]
    path = row.get("storage_path")
    _validate_storage_path(path)
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="File not found on disk")

    return FileResponse(
        path=path,
        filename=row.get("original_filename"),
        media_type=row.get("content_type"),
    )


@router.delete("/{document_id}")
async def delete_document(
    document_id: int,
    ctx: dict = Depends(require_role("manager")),
):
    result = await bridge.get(
        bridge.db_service, f"/documents/{document_id}",
        params={
            "org_id": ctx["org_id"],
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    if not result:
        raise HTTPException(status_code=404, detail="Document not found or not yours to delete")

    doc = result[0]
    path = doc.get("storage_path")
    try:
        _validate_storage_path(path)
        if os.path.exists(path):
            os.remove(path)
    except Exception:
        logger.warning("Failed to delete file from disk: %s", path)

    await bridge.delete(bridge.db_service, f"/documents/{document_id}/summaries")
    await bridge.delete(bridge.db_service, f"/documents/{document_id}")
    return {"ok": True}


class SummarizeRequest(BaseModel):
    provider: str = "openai"
    api_key: str = ""
    model: str = "gpt-4o-mini"

    @field_validator("provider")
    @classmethod
    def validate_provider(cls, v: str) -> str:
        v = v.strip().lower()
        if v not in settings.ALLOWED_PROVIDERS:
            raise ValueError(f"Unsupported provider: {v}. Allowed: {', '.join(sorted(settings.ALLOWED_PROVIDERS))}")
        return v

    @field_validator("model")
    @classmethod
    def validate_model(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("Model name is required")
        if len(v) > 200:
            raise ValueError("Model name too long")
        return v


@router.post("/{document_id}/summarize")
async def summarize_document(
    document_id: int,
    req: SummarizeRequest,
    ctx: dict = Depends(require_role("member")),
):
    user_id = ctx["user_id"]
    org_id = ctx["org_id"]

    result = await bridge.get(
        bridge.db_service, f"/documents/{document_id}/text",
        params={
            "org_id": org_id,
            "user_id": ctx["user_id"],
            "is_admin": ctx["is_admin"],
            "is_manager": ctx["is_manager"],
        },
    )
    if not result:
        raise HTTPException(status_code=404, detail="Document not found")

    doc = result[0]
    text_content = doc.get("extracted_text") or ""
    if not text_content.strip():
        raise HTTPException(status_code=422, detail="No extractable text found in this document")

    truncated = text_content[:15000]
    system_prompt = (
        "You are a document summarization assistant for StratRoom, an enterprise governance platform. "
        "Provide a concise, structured summary of the document. Include: "
        "1) Document type and purpose, "
        "2) Key points and findings, "
        "3) Action items or recommendations if any. "
        "Format with clear headers and bullet points."
    )

    try:
        summary = await call_llm(
            provider=req.provider,
            api_key=req.api_key,
            model=req.model,
            system_prompt=system_prompt,
            messages=[{"role": "user", "content": f"Summarize this document ({doc.get('original_filename', '')}):\n\n{truncated}"}],
            max_tokens=2048,
        )
    except Exception:
        logger.exception("Document summarization failed [doc=%s user=%s]", document_id, user_id)
        raise HTTPException(status_code=502, detail="Summarization failed")

    await bridge.post(bridge.db_service, f"/documents/{document_id}/summarize", json={
        "summary": summary[:10000],
        "provider": req.provider,
        "model": req.model,
    })

    return {"summary": summary, "provider": req.provider, "model": req.model}
